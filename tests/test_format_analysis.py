import base64
from io import BytesIO

import docx
import pytest
from pypdf import PdfWriter

from app.core.parsing.format_analysis import (
    _looks_multi_column, analyze_docx_format, analyze_format, analyze_pdf_format,
)

# A minimal, valid 1x1 transparent PNG -- small enough to inline, real
# enough for python-docx's add_picture() to accept as a genuine image.
_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


# --- pure multi-column heuristic (no file I/O needed) -------------------

def test_wide_midline_gaps_flag_as_multi_column():
    layout_text = "\n".join([
        "Python programming          Led backend migration to microservices",
        "AWS and Docker              Reduced latency by 40 percent overall",
        "SQL and NoSQL databases     Managed a team of four engineers here",
    ])
    assert _looks_multi_column(layout_text) is True


def test_normal_single_column_text_is_not_flagged():
    layout_text = "\n".join([
        "Experienced backend engineer with a focus on distributed systems.",
        "Led the migration of the ingestion pipeline to a queue-based design.",
        "Reduced end-to-end latency by 40% across the primary service.",
    ])
    assert _looks_multi_column(layout_text) is False


def test_empty_text_is_not_flagged():
    assert _looks_multi_column("") is False


# --- DOCX format analysis, using real generated documents ---------------

def _docx_bytes(build_fn) -> bytes:
    document = docx.Document()
    build_fn(document)
    buf = BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_docx_with_table_is_flagged():
    def build(doc):
        doc.add_paragraph("Jordan Alvarez -- jordan@example.com")
        table = doc.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Python"
        table.rows[0].cells[1].text = "AWS"

    result = analyze_docx_format(_docx_bytes(build))
    assert result.has_tables is True
    assert any("table" in f.lower() for f in result.flags)


def test_docx_without_table_is_not_flagged():
    def build(doc):
        doc.add_paragraph("Jordan Alvarez -- jordan@example.com")
        doc.add_paragraph("Backend engineer with five years of experience.")

    result = analyze_docx_format(_docx_bytes(build))
    assert result.has_tables is False


def test_docx_embedded_image_is_detected():
    import tempfile

    def build(doc):
        doc.add_paragraph("Jordan Alvarez")
        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp.write(_TINY_PNG)
            tmp.flush()
            doc.add_picture(tmp.name)

    result = analyze_docx_format(_docx_bytes(build))
    assert result.embedded_image_count >= 1
    assert any("image" in f.lower() for f in result.flags)


def test_docx_contact_only_in_header_is_flagged():
    def build(doc):
        doc.sections[0].header.paragraphs[0].text = "jordan@example.com | 555-123-4567"
        doc.add_paragraph("Backend engineer with five years of experience building APIs.")

    result = analyze_docx_format(_docx_bytes(build))
    assert result.contact_info_only_in_header_footer is True


def test_docx_contact_in_body_is_not_flagged():
    def build(doc):
        doc.add_paragraph("Jordan Alvarez -- jordan@example.com -- 555-123-4567")
        doc.add_paragraph("Backend engineer.")

    result = analyze_docx_format(_docx_bytes(build))
    assert result.contact_info_only_in_header_footer is False


def test_docx_risk_count_reflects_flags_found():
    def build(doc):
        table = doc.add_table(rows=1, cols=1)
        table.rows[0].cells[0].text = "Python"

    result = analyze_docx_format(_docx_bytes(build))
    assert result.risk_count >= 1


# --- PDF format analysis --------------------------------------------------

def test_pdf_blank_page_is_flagged_as_scanned():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = BytesIO()
    writer.write(buf)

    result = analyze_pdf_format(buf.getvalue())
    assert result.file_type == "pdf_scanned"
    assert any("scanned" in f.lower() for f in result.flags)


def test_pdf_unopenable_bytes_return_unknown_gracefully():
    result = analyze_pdf_format(b"not a real pdf")
    assert result.file_type == "unknown"
    assert result.flags


# --- dispatch by extension -----------------------------------------------

def test_analyze_format_dispatches_by_extension():
    def build(doc):
        doc.add_paragraph("Some content.")

    result = analyze_format(_docx_bytes(build), "resume.docx")
    assert result.file_type == "docx"


def test_analyze_format_unknown_extension():
    result = analyze_format(b"whatever", "resume.txt")
    assert result.file_type == "unknown"
