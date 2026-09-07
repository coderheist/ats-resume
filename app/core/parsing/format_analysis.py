"""
Real ATS format/layout analysis on the raw uploaded file -- resolves the
limitation stated in scoring/screening_report.py's docstring: true layout
analysis needs the original file's bytes, which aren't available once a
resume has already become a structured JsonResume. This module runs at
the one place those bytes still exist -- alongside file_extraction.py, on
the same PDF/DOCX upload -- before anything gets thrown away.

Deliberately scoped to what's genuinely detectable without the GB-scale
CV/layout models parser_interface.py's Docling/Marker adapters use (see
that module's docstring): pypdf and python-docx expose enough structural
information (embedded images, tables, header/footer text, and PDF's
layout-preserving text-extraction mode) to catch the specific failure
modes the spec calls out, without needing those heavier dependencies.

What's still out of scope, honestly: OCR for scanned/image-only PDFs
(file_extraction.py already flags these separately), and precise column
*boundaries* (the multi-column check below is a real heuristic, not a
guarantee -- a false positive/negative is possible on unusual layouts).
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from io import BytesIO

_EMAIL_PATTERN = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
_PHONE_PATTERN = re.compile(r"(?:\+?\d[\d\-\s()]{8,}\d)")


@dataclass
class FormatAnalysis:
    file_type: str  # "pdf_text" | "pdf_scanned" | "docx" | "unknown"
    likely_multi_column: bool = False
    embedded_image_count: int = 0
    has_tables: bool = False
    contact_info_only_in_header_footer: bool = False
    flags: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "file_type": self.file_type,
            "likely_multi_column": self.likely_multi_column,
            "embedded_image_count": self.embedded_image_count,
            "has_tables": self.has_tables,
            "contact_info_only_in_header_footer": self.contact_info_only_in_header_footer,
            "flags": self.flags,
        }

    @property
    def risk_count(self) -> int:
        """How many distinct format risks were found -- used by
        screening_report.py to fold this into the ATS Readability score
        without needing to know the specific flag types."""
        return sum([
            self.file_type == "pdf_scanned",
            self.likely_multi_column,
            self.has_tables,
            self.embedded_image_count > 0,
            self.contact_info_only_in_header_footer,
        ])


def _looks_multi_column(layout_text: str) -> bool:
    """Pure heuristic, independently testable: a large contiguous gap in
    the *middle* of a line (not just trailing whitespace) is what two
    side-by-side text blocks look like once flattened to plain text via
    a layout-preserving extraction. Real signal, not a guarantee -- an
    unusual single-column layout with wide internal spacing could
    false-positive here."""
    lines = [line for line in layout_text.split("\n") if line.strip()]
    if not lines:
        return False
    suspect_lines = 0
    for line in lines:
        stripped = line.rstrip()
        if not stripped:
            continue
        match = re.search(r"\S(\s{6,})\S", stripped)
        if match and match.start() > 5 and match.end() < len(stripped) - 5:
            suspect_lines += 1
    return (suspect_lines / len(lines)) > 0.15


def analyze_pdf_format(file_bytes: bytes) -> FormatAnalysis:
    from pypdf import PdfReader

    try:
        reader = PdfReader(BytesIO(file_bytes))
    except Exception:
        return FormatAnalysis(file_type="unknown", flags=["Couldn't open this PDF for format analysis."])

    if reader.is_encrypted:
        # See file_extraction.py's _extract_pdf for the full explanation:
        # decrypt() returns 0 on failure rather than raising, so it must be
        # checked explicitly -- not just wrapped in try/except -- or a
        # non-blank-password PDF silently proceeds fully encrypted and
        # crashes the next line that touches reader.pages.
        if not reader.decrypt(""):
            return FormatAnalysis(file_type="unknown", flags=["This PDF is password-protected -- format analysis skipped."])

    try:
        plain_text = "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    except Exception:
        return FormatAnalysis(file_type="unknown", flags=["Couldn't extract text from this PDF for format analysis."])
    if not plain_text:
        return FormatAnalysis(
            file_type="pdf_scanned",
            flags=["No extractable text -- this looks like a scanned/image-only PDF. "
                   "ATS systems cannot reliably parse this without OCR."],
        )

    analysis = FormatAnalysis(file_type="pdf_text")

    # Embedded images -- icons, graphic skill bars, decorative elements,
    # or (most seriously) real content baked into an image instead of
    # actual text -- via pypdf's per-page embedded-image listing.
    image_count = 0
    for page in reader.pages:
        try:
            image_count += len(page.images)
        except Exception:  # noqa: BLE001 -- a single bad page shouldn't fail the whole scan
            pass
    analysis.embedded_image_count = image_count
    if image_count > 0:
        analysis.flags.append(
            f"{image_count} embedded image(s) detected -- icons, graphic skill bars, or decorative "
            f"elements can confuse ATS text extraction, and any real content placed inside an image "
            f"(rather than as actual text) will not be seen by an ATS at all."
        )

    # Multi-column heuristic: pypdf's layout-preserving extraction mode
    # keeps the PDF's original horizontal spacing. A large contiguous gap
    # in the *middle* of a line (not just trailing whitespace) is what
    # two side-by-side text blocks look like once flattened to plain
    # text -- a real signal, not a guess, but still a heuristic: unusual
    # single-column layouts with wide internal spacing could false-
    # positive here.
    try:
        layout_text = "\n".join(
            (page.extract_text(extraction_mode="layout") or "") for page in reader.pages
        )
        if _looks_multi_column(layout_text):
            analysis.likely_multi_column = True
            analysis.flags.append(
                "This document's layout may use multiple columns. Many ATS parsers read left-to-right "
                "across the whole page and can interleave column content into garbled, out-of-order text "
                "-- a single-column layout is safest."
            )
    except TypeError:
        # Older pypdf without extraction_mode support -- skip rather than
        # fail the whole analysis over one unavailable check.
        pass

    return analysis


def analyze_docx_format(file_bytes: bytes) -> FormatAnalysis:
    import docx

    try:
        document = docx.Document(BytesIO(file_bytes))
    except Exception:
        return FormatAnalysis(file_type="unknown", flags=["Couldn't open this DOCX for format analysis."])

    analysis = FormatAnalysis(file_type="docx")

    if document.tables:
        analysis.has_tables = True
        analysis.flags.append(
            f"{len(document.tables)} table(s) detected -- tables are a common source of ATS parsing "
            f"failures (content can be read out of order or dropped entirely). A single-column layout "
            f"without tables is safest."
        )

    image_count = len(document.inline_shapes)
    analysis.embedded_image_count = image_count
    if image_count > 0:
        analysis.flags.append(
            f"{image_count} embedded image(s) detected -- any real content placed inside an image "
            f"rather than as actual text will not be seen by an ATS at all."
        )

    body_text = "\n".join(p.text for p in document.paragraphs)
    header_footer_text = ""
    for section in document.sections:
        header_footer_text += "\n" + "\n".join(p.text for p in section.header.paragraphs)
        header_footer_text += "\n" + "\n".join(p.text for p in section.footer.paragraphs)

    body_has_contact = bool(_EMAIL_PATTERN.search(body_text) or _PHONE_PATTERN.search(body_text))
    header_footer_has_contact = bool(
        _EMAIL_PATTERN.search(header_footer_text) or _PHONE_PATTERN.search(header_footer_text)
    )
    if header_footer_has_contact and not body_has_contact:
        analysis.contact_info_only_in_header_footer = True
        analysis.flags.append(
            "Contact information appears to be only in the header/footer -- many ATS parsers skip "
            "headers and footers entirely. Put your email and phone number in the main document body too."
        )

    return analysis


def analyze_format(file_bytes: bytes, filename: str) -> FormatAnalysis:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return analyze_pdf_format(file_bytes)
    if lower.endswith(".docx"):
        return analyze_docx_format(file_bytes)
    return FormatAnalysis(file_type="unknown", flags=["Format analysis only runs on .pdf/.docx uploads."])
